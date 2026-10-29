"""Energiföretagen weekly statistics: an independent check of the reservoir data.

Energiföretagen, the Swedish energy industry association, publishes one workbook per year
("Veckostatistik") with, for each ISO week, hydro production ("Vattenkraft"), inflow
("Tillrinning i GWh, ej spillkorrigerad") and the reservoir filling at the end of the week
("Magasinsfyllnad i %", of the capacity printed under it). Comparing them with the ENTSO-E series
shows which ENTSO-E label corresponds to each week and which labelling convention reproduces the
published inflow. The workbooks are downloaded on demand into ``data/external/`` and are never
committed; only the comparison statistics are.
"""

from __future__ import annotations

import hashlib
import re
from collections.abc import Callable, Mapping, Sequence
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

import numpy as np
import openpyxl
import pandas as pd
import requests

from watervalues.config import PROJECT_ROOT, AppConfig
from watervalues.data.aggregate import hourly_energy
from watervalues.data.build import unfolded_hydro, week_label
from watervalues.data.cache import read_dataset
from watervalues.data.calendar import iso_week_start
from watervalues.data.reservoir import WEEK, LabelSemantics, implied_inflow, labelled_series

ROW_LABELS = {"Vecka": "week", "Datum": "monday", "Vattenkraft": "hydro_gwh"}
ROW_PREFIXES = {"Tillrinning": "inflow_gwh", "Magasinsfyllnad": "fill_pct"}
VALUE_COLUMNS = ("hydro_gwh", "inflow_gwh", "fill_pct")
OFFSETS = (-1, 0, 1)  # the ENTSO-E label w + k is compared with workbook week w
_WEEK = re.compile(r"(\d{4})-(\d{2})")
_FIRST_VALUE_COLUMN = 2  # row labels sit in the first two columns, weeks from the third on
_PERCENT = 100.0
_TAIL_PERCENTILE = 90
_XLSX_MAGIC = b"PK"  # an xlsx workbook is a zip archive
_SHORT_DIGEST = 16
_MIN_PAIRS = 3  # fewer pairs give no meaningful correlation


class WorkbookError(ValueError):
    """A workbook is missing or does not have the expected layout."""


@dataclass(frozen=True)
class Validation:
    """Energiföretagen workbooks compared with the ENTSO-E series (all energies in GWh)."""

    first_week: str
    last_week: str
    not_reported: list[str]  # weeks with a zero filling or inflow (not filled in yet)
    entry_errors: list[str]  # weeks with a filling above 100 % (e.g. a lost decimal point)
    storage_offsets: pd.DataFrame  # per offset k: median and 90th percentile of |diff|, n
    storage_by_year: pd.DataFrame  # offset 0: median and max |diff| per ISO year
    hydro_lag_corr: dict[int, float]  # correlation with ENTSO-E hydro of week w + k
    hydro_ratio_by_year: pd.Series  # ENTSO-E / Energiföretagen hydro production
    inflow_fit: pd.DataFrame  # per convention and window: corr, RMSE, bias, n
    inflow_rmse_by_year: pd.DataFrame  # per ISO year: RMSE of each convention and window


def download_workbooks(
    urls: Mapping[int, str],
    folder: Path,
    timeout_s: float,
    *,
    force: bool = False,
    get: Callable[..., requests.Response] | None = None,
) -> list[Path]:
    """Download the yearly workbooks that are not cached yet.

    Parameters
    ----------
    urls : mapping of int to str
        Workbook URL per year.
    folder : Path
        Target directory, created if needed.
    timeout_s : float
        Timeout per request [s].
    force : bool
        Download again even if the file exists.
    get : callable or None
        ``requests.get``-like function (injected in tests).

    Returns
    -------
    list of Path
        Local workbooks in year order.

    Raises
    ------
    WorkbookError
        If a response is not an xlsx file.
    """
    fetch = requests.get if get is None else get
    folder.mkdir(parents=True, exist_ok=True)
    paths = []
    for year in sorted(urls):
        path = folder / Path(urlparse(urls[year]).path).name
        if force or not path.exists():
            response = fetch(urls[year], timeout=timeout_s)
            response.raise_for_status()
            if not response.content.startswith(_XLSX_MAGIC):
                raise WorkbookError(f"{urls[year]} did not return an xlsx workbook")
            tmp = path.with_suffix(".tmp")
            tmp.write_bytes(response.content)
            tmp.replace(path)
        paths.append(path)
    return paths


def _row_key(label: object) -> str | None:
    if not isinstance(label, str):
        return None
    text = label.strip()
    if text in ROW_LABELS:
        return ROW_LABELS[text]
    return next((name for prefix, name in ROW_PREFIXES.items() if text.startswith(prefix)), None)


def _number(values: Sequence[object], i: int) -> float:
    value = values[i] if i < len(values) else None
    return float(value) if isinstance(value, int | float) else np.nan


def read_workbook(path: Path, iso_year: int | None = None) -> pd.DataFrame:
    """Weekly hydro production, inflow and end-of-week filling from one workbook.

    Parameters
    ----------
    path : Path
        A yearly "Veckostatistik" workbook.
    iso_year : int or None
        Keep only the weeks of this ISO year (a workbook may carry an empty column for the
        first week of the next year).

    Returns
    -------
    pandas.DataFrame
        ``week_start`` (Monday 00:00 UTC), ``iso_year``, ``iso_week``, ``hydro_gwh`` and
        ``inflow_gwh`` [GWh] and ``fill_pct`` [%], one row per ISO week. Annual totals and
        other summary columns are dropped.

    Raises
    ------
    WorkbookError
        If a row is missing or a week is not dated on the Monday that starts it.
    """
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows: dict[str, list[object]] = {}
        for row in workbook.worksheets[0].iter_rows(values_only=True):
            for label in row[:_FIRST_VALUE_COLUMN]:
                key = _row_key(label)
                if key is not None and key not in rows:
                    rows[key] = list(row[_FIRST_VALUE_COLUMN:])
    finally:
        workbook.close()
    missing = sorted((set(ROW_LABELS.values()) | set(ROW_PREFIXES.values())) - set(rows))
    if missing:
        raise WorkbookError(f"{path.name}: rows not found: {missing}")
    records = []
    for i, week in enumerate(rows["week"]):
        match = _WEEK.fullmatch(str(week).strip()) if week is not None else None
        monday = rows["monday"][i] if i < len(rows["monday"]) else None
        if match is None or not isinstance(monday, datetime):
            continue  # annual totals and other summary columns
        year, iso_week = int(match[1]), int(match[2])
        if iso_year is not None and year != iso_year:
            continue
        start = iso_week_start(year, iso_week)
        if pd.Timestamp(monday).tz_localize("UTC") != start:
            raise WorkbookError(f"{path.name}: week {week} is dated {monday:%Y-%m-%d}")
        values = {column: _number(rows[column], i) for column in VALUE_COLUMNS}
        records.append({"week_start": start, "iso_year": year, "iso_week": iso_week, **values})
    return pd.DataFrame(records)


def _fit(reference: np.ndarray, estimate: np.ndarray) -> dict[str, float]:
    ok = ~np.isnan(reference) & ~np.isnan(estimate)
    a, b = reference[ok], estimate[ok]
    if len(a) < _MIN_PAIRS:
        return {"corr": np.nan, "rmse_gwh": np.nan, "bias_gwh": np.nan, "n": len(a)}
    return {
        "corr": float(np.corrcoef(a, b)[0, 1]),
        "rmse_gwh": float(np.sqrt(np.mean((b - a) ** 2))),
        "bias_gwh": float(np.mean(b - a)),
        "n": int(ok.sum()),
    }


def compare(
    workbooks: pd.DataFrame,
    published_gwh: pd.Series,
    hydro_gwh: pd.Series,
    max_storage_gwh: float,
    winter_weeks: tuple[int, int],
) -> Validation:
    """Compare the workbooks with the ENTSO-E reservoir labels and hydro production.

    Parameters
    ----------
    workbooks : pandas.DataFrame
        Output of :func:`read_workbook`, one or more years.
    published_gwh : pandas.Series
        ENTSO-E reservoir values [GWh] indexed by the start of their labelled ISO week.
    hydro_gwh : pandas.Series
        ENTSO-E hydro production [GWh] per ISO week (week 53 separate), indexed by week start.
    max_storage_gwh : float
        Capacity that the workbooks' filling percentages refer to [GWh].
    winter_weeks : tuple of int
        Inclusive range of ISO weeks of the winter window.

    Returns
    -------
    Validation
        Comparison statistics.
    """
    ef = workbooks.set_index("week_start").sort_index()
    t = pd.DatetimeIndex(ef.index)
    years = ef["iso_year"].to_numpy()
    # The Swedish aggregate is never empty and never without inflow: zeros mean "not reported".
    zero = (ef["fill_pct"] == 0) | (ef["inflow_gwh"] == 0)
    too_full = ef["fill_pct"] > _PERCENT
    fill = ef["fill_pct"].where((ef["fill_pct"] != 0) & ~too_full)
    ef["inflow_gwh"] = ef["inflow_gwh"].where(ef["inflow_gwh"] != 0)
    storage = (fill / _PERCENT * max_storage_gwh).to_numpy()

    offsets = {}
    for k in OFFSETS:
        diff = np.abs(storage - published_gwh.reindex(t + k * WEEK).to_numpy())
        ok = ~np.isnan(diff)
        offsets[k] = {
            "median_abs_gwh": float(np.median(diff[ok])),
            f"p{_TAIL_PERCENTILE}_abs_gwh": float(np.percentile(diff[ok], _TAIL_PERCENTILE)),
            "n": int(ok.sum()),
        }
    same_week = pd.Series(np.abs(storage - published_gwh.reindex(t).to_numpy()), index=ef.index)
    by_year = same_week.groupby(years).agg(["median", "max"])
    by_year.columns = ["median_abs_gwh", "max_abs_gwh"]

    hydro_ef = ef["hydro_gwh"].to_numpy()
    lag_corr = {
        k: _fit(hydro_ef, hydro_gwh.reindex(t + k * WEEK).to_numpy())["corr"] for k in OFFSETS
    }
    hydro_entsoe = hydro_gwh.reindex(t).to_numpy()
    both = ~np.isnan(hydro_ef) & ~np.isnan(hydro_entsoe)
    ratio = (
        pd.Series(hydro_entsoe[both]).groupby(years[both]).sum()
        / pd.Series(hydro_ef[both]).groupby(years[both]).sum()
    )

    inflow_ef = ef["inflow_gwh"].to_numpy()
    winter = (ef["iso_week"] >= winter_weeks[0]).to_numpy() & (
        ef["iso_week"] <= winter_weeks[1]
    ).to_numpy()
    windows = {"all weeks": np.ones(len(ef), dtype=bool), "winter weeks": winter}
    fits, per_year = [], {}
    for semantics in LabelSemantics:
        implied = implied_inflow(published_gwh, hydro_gwh, semantics).reindex(t).to_numpy()
        for window, mask in windows.items():
            fits.append(
                {
                    "convention": semantics.value,
                    "window": window,
                    **_fit(inflow_ef[mask], implied[mask]),
                }
            )
            per_year[f"{semantics.value} ({window})"] = {
                int(y): _fit(inflow_ef[mask & (years == y)], implied[mask & (years == y)])[
                    "rmse_gwh"
                ]
                for y in np.unique(years)
            }
    labels = [week_label(y, w) for y, w in zip(ef["iso_year"], ef["iso_week"], strict=True)]
    return Validation(
        first_week=labels[0],
        last_week=labels[-1],
        not_reported=[label for label, bad in zip(labels, zero, strict=True) if bad],
        entry_errors=[label for label, bad in zip(labels, too_full, strict=True) if bad],
        storage_offsets=pd.DataFrame(offsets).T.astype({"n": int}).rename_axis("k"),
        storage_by_year=by_year.rename_axis("iso_year"),
        hydro_lag_corr=lag_corr,
        hydro_ratio_by_year=ratio.rename_axis("iso_year"),
        inflow_fit=pd.DataFrame(fits).set_index(["convention", "window"]),
        inflow_rmse_by_year=pd.DataFrame(per_year).rename_axis("iso_year"),
    )


def validation_report(
    result: Validation, files: Sequence[tuple[str, str]], source: str
) -> list[str]:
    """Human-readable comparison, with the workbooks' names and SHA-256 digests."""
    lines = [
        f"Energiföretagen weekly statistics vs ENTSO-E ({source}), "
        f"ISO weeks {result.first_week} to {result.last_week}",
        "Workbooks: " + ", ".join(f"{name} ({digest[:_SHORT_DIGEST]})" for name, digest in files),
        f"Weeks with a zero filling or inflow (not reported), ignored: {result.not_reported}",
        f"Weeks with a filling above 100 % (entry errors), ignored: {result.entry_errors}",
        "\n1) Filling at the end of workbook week w vs the ENTSO-E value labelled w + k [GWh]",
        result.storage_offsets.round(1).to_string(),
        "   offset 0, per ISO year:",
        result.storage_by_year.round(1).to_string(),
        "\n2) Hydro production",
        "   correlation with ENTSO-E week w + k: "
        + ", ".join(f"k={k:+d}: {c:.3f}" for k, c in result.hydro_lag_corr.items()),
        "   ENTSO-E / Energiföretagen per ISO year: "
        + ", ".join(f"{y}: {r:.3f}" for y, r in result.hydro_ratio_by_year.items()),
        "\n3) Workbook inflow vs inflow implied by the ENTSO-E labels [GWh/week]",
        result.inflow_fit.round(3).to_string(),
        "\n4) RMSE of the implied inflow per ISO year [GWh/week]",
        result.inflow_rmse_by_year.round(0).to_string(),
    ]
    return lines


def validate_energiforetagen(
    cfg: AppConfig,
    source: str,
    *,
    force: bool = False,
    get: Callable[..., requests.Response] | None = None,
) -> tuple[Validation, list[str]]:
    """Download the workbooks (if needed) and compare them with the cached ENTSO-E data.

    Parameters
    ----------
    cfg : AppConfig
        Configuration with ``data.energiforetagen``.
    source : str
        Raw-cache source of the ENTSO-E data.
    force : bool
        Download the workbooks again.
    get : callable or None
        ``requests.get``-like function (injected in tests).

    Returns
    -------
    tuple of (Validation, list of str)
        Statistics and the report lines.
    """
    settings = cfg.data.energiforetagen
    paths = download_workbooks(
        settings.workbooks, PROJECT_ROOT / settings.dir, settings.timeout_s, force=force, get=get
    )
    workbooks = pd.concat(
        [
            read_workbook(path, year)
            for year, path in zip(sorted(settings.workbooks), paths, strict=True)
        ],
        ignore_index=True,
    )
    if workbooks["week_start"].duplicated().any():
        raise WorkbookError("the workbooks contain the same week more than once")
    raw_dir = PROJECT_ROOT / cfg.data.raw_dir
    published = labelled_series(read_dataset(raw_dir, source, "reservoir", cfg.data.reservoir.area))
    generation = read_dataset(raw_dir, source, "generation", cfg.data.country)
    hydro = unfolded_hydro(cfg, hourly_energy(generation, "value_mw", by=("psr_type",)))
    result = compare(
        workbooks,
        published,
        hydro,
        settings.max_storage_gwh,
        tuple(cfg.data.reservoir.winter_weeks),
    )
    files = [(path.name, hashlib.sha256(path.read_bytes()).hexdigest()) for path in paths]
    return result, validation_report(result, files, source)
