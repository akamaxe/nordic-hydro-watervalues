"""Tests for the Energiföretagen workbook reader and comparison (synthetic workbooks, offline)."""

import io
from datetime import datetime

import numpy as np
import openpyxl
import pandas as pd
import pytest

from watervalues.config import load_config
from watervalues.data.cache import write_dataset
from watervalues.data.energiforetagen import (
    WorkbookError,
    compare,
    download_workbooks,
    read_workbook,
    validate_energiforetagen,
    validation_report,
)

MAX_GWH = 33_748.0
WINTER = (2, 14)


def workbook_bytes(weeks, hydro, inflow, fill, mondays=None, rows=None) -> bytes:
    """A sheet laid out like Energiföretagen's: labels in columns A-B, one column per week."""
    labels = [f"{y}-{w:02d}" for y, w in weeks]
    dates = mondays or [datetime.fromisocalendar(y, w, 1) for y, w in weeks]
    layout = {
        3: [None, "Vecka", *labels, labels[0]],  # last column: annual total, no date
        4: [None, "Datum", *dates, labels[-1]],
        5: [None, "Version: A", *(["GWh"] * len(weeks))],
        9: ["Vattenkraft", None, *hydro, sum(hydro)],
        58: ["Tillrinning i GWh", None, *inflow, sum(inflow)],
        59: ["(ej spillkorrigerad)", None],
        60: ["Magasinsfyllnad i %", None, *fill],
        61: ["(max mag 33 748 GWh)", None],
    }
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row, values in layout.items():
        if rows is not None and row not in rows:
            continue
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row, column=column, value=value)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_read_workbook_keeps_the_weeks_of_its_year(tmp_path) -> None:
    weeks = [(2019, 1), (2019, 2), (2019, 3), (2020, 1)]  # 2020-01: empty placeholder column
    path = tmp_path / "vecka_2019.xlsx"
    path.write_bytes(
        workbook_bytes(
            weeks, [1172.5, 1319.3, 1453.4, 0], [558, 592.7, 394.9, 0], [53.5, 51.4, 48.2, 0]
        )
    )
    frame = read_workbook(path, 2019)
    assert frame["iso_week"].tolist() == [1, 2, 3]
    assert frame["week_start"].iloc[0] == pd.Timestamp("2018-12-31", tz="UTC")
    assert frame["fill_pct"].tolist() == [53.5, 51.4, 48.2]
    assert len(read_workbook(path)) == 4


def test_misdated_week_and_missing_rows_are_rejected(tmp_path) -> None:
    path = tmp_path / "bad.xlsx"
    tuesday = datetime.fromisocalendar(2019, 2, 2)
    path.write_bytes(workbook_bytes([(2019, 2)], [1.0], [1.0], [1.0], mondays=[tuesday]))
    with pytest.raises(WorkbookError, match="dated 2019-01-08"):
        read_workbook(path)
    path.write_bytes(workbook_bytes([(2019, 2)], [1.0], [1.0], [1.0], rows={3, 4, 9, 60}))
    with pytest.raises(WorkbookError, match="inflow_gwh"):
        read_workbook(path)


class FakeResponse:
    def __init__(self, content: bytes) -> None:
        self.content = content

    def raise_for_status(self) -> None:
        pass


def test_download_skips_cached_files_and_rejects_other_content(tmp_path) -> None:
    calls = []

    def get(url, timeout):
        calls.append(url)
        return FakeResponse(b"PK\x03\x04 workbook")

    urls = {2019: "https://example.org/a/vecka_2019.xlsx"}
    paths = download_workbooks(urls, tmp_path, 5, get=get)
    assert paths == [tmp_path / "vecka_2019.xlsx"] and paths[0].read_bytes().startswith(b"PK")
    download_workbooks(urls, tmp_path, 5, get=get)
    assert len(calls) == 1
    with pytest.raises(WorkbookError, match="xlsx"):
        download_workbooks(
            urls, tmp_path, 5, force=True, get=lambda url, timeout: FakeResponse(b"<html>")
        )


@pytest.fixture
def reservoir():
    """End-of-week storage, production and inflow for 2018-W52 .. 2020-W53 [GWh]."""
    rng = np.random.default_rng(3)
    starts = pd.date_range("2018-12-24", "2020-12-28", freq="7D", tz="UTC")
    week = np.asarray(starts.isocalendar()["week"])
    inflow = 400 + 2500 * np.exp(-(((week - 24) / 4) ** 2)) + rng.normal(0, 10, len(starts))
    hydro = inflow.mean() + rng.normal(0, 100, len(starts))
    end_storage = 15_000 + np.cumsum(inflow - hydro)
    return starts, hydro, inflow, end_storage


def workbook_frame(starts, hydro, inflow, end_storage) -> pd.DataFrame:
    iso = starts.isocalendar()
    return pd.DataFrame(
        {
            "week_start": starts,
            "iso_year": np.asarray(iso["year"]),
            "iso_week": np.asarray(iso["week"]),
            "hydro_gwh": hydro,
            "inflow_gwh": inflow,
            "fill_pct": np.round(end_storage / MAX_GWH * 100, 1),
        }
    )


def test_compare_finds_end_of_week_labels_and_skips_bad_weeks(reservoir) -> None:
    starts, hydro, inflow, end_storage = reservoir
    published = pd.Series(end_storage, index=starts)
    ef = workbook_frame(starts[1:], hydro[1:], inflow[1:], end_storage[1:])
    ef.loc[10, "fill_pct"] = 7710.0  # a lost decimal point
    ef.loc[20, ["fill_pct", "inflow_gwh"]] = 0.0  # not reported yet
    result = compare(ef, published, pd.Series(hydro, index=starts), MAX_GWH, WINTER)
    assert result.entry_errors == ["2019-W11"]
    assert result.not_reported == ["2019-W21"]
    offsets = result.storage_offsets["median_abs_gwh"]
    assert offsets.loc[0] < 0.05 / 100 * MAX_GWH  # rounding of the printed filling
    assert min(offsets.loc[-1], offsets.loc[1]) > 10 * offsets.loc[0]
    np.testing.assert_allclose(result.storage_by_year["implied_capacity_gwh"], MAX_GWH, rtol=1e-3)
    fit = result.inflow_fit["rmse_gwh"]
    assert fit.loc[("end_of_week", "all weeks")] == pytest.approx(0.0, abs=1e-6)
    assert fit.loc[("weekly_mean", "all weeks")] > 50
    assert result.hydro_lag_corr[0] == pytest.approx(1.0)
    lines = validation_report(result, [("vecka_2019.xlsx", "0" * 64)], "test")
    assert any("entry errors" in line and "2019-W11" in line for line in lines)


def test_validate_reads_each_workbook_for_its_own_year(reservoir, tmp_path) -> None:
    starts, hydro, inflow, end_storage = reservoir
    raw = tmp_path / "raw"
    iso = starts.isocalendar()
    labels = pd.DataFrame(
        {
            "iso_year": np.asarray(iso["year"]),
            "iso_week": np.asarray(iso["week"]),
            "stored_mwh": end_storage * 1000.0,
        }
    )
    write_dataset(labels, raw, "filelib", "reservoir", "SE", labels["iso_year"])
    hours = pd.date_range(starts[1], starts[-1] + pd.Timedelta(days=7), freq="h", inclusive="left")
    hydro_mw = np.repeat(hydro[1:], 168) * 1000 / 168
    generation = pd.DataFrame(
        {
            "start_utc": hours,
            "resolution_min": 60,
            "psr_type": "B12",
            "value_mw": hydro_mw,
            "consumption_mw": 0.0,
        }
    )
    write_dataset(generation, raw, "filelib", "generation", "SE", hours.year)

    ef = workbook_frame(starts[1:], hydro[1:], inflow[1:], end_storage[1:])
    served = {}
    for year in (2019, 2020):
        part = ef[ef["iso_year"] == year]
        weeks = list(zip(part["iso_year"], part["iso_week"], strict=True))
        hydro_w, inflow_w, fill_w = (
            part[c].tolist() for c in ("hydro_gwh", "inflow_gwh", "fill_pct")
        )
        if year == 2019:  # the next year's first week as an empty column
            weeks, hydro_w, inflow_w, fill_w = (
                [*weeks, (2020, 1)],
                [*hydro_w, 0],
                [*inflow_w, 0],
                [*fill_w, 0],
            )
        served[f"https://example.org/vecka_{year}.xlsx"] = workbook_bytes(
            weeks, hydro_w, inflow_w, fill_w
        )

    cfg = load_config()
    settings = cfg.data.energiforetagen.model_copy(
        update={
            "dir": tmp_path / "ef",
            "workbooks": {y: f"https://example.org/vecka_{y}.xlsx" for y in (2019, 2020)},
        }
    )
    data = cfg.data.model_copy(update={"raw_dir": raw, "energiforetagen": settings})
    period = cfg.period.model_copy(update={"first_iso_year": 2019, "last_iso_year": 2020})
    cfg = cfg.model_copy(update={"data": data, "period": period})

    result, lines = validate_energiforetagen(
        cfg, "filelib", get=lambda url, timeout: FakeResponse(served[url])
    )
    assert (result.first_week, result.last_week) == ("2019-W01", "2020-W53")
    assert result.not_reported == []
    assert result.inflow_fit.loc[("end_of_week", "all weeks"), "rmse_gwh"] == pytest.approx(
        0, abs=1e-6
    )
    assert lines[1].startswith("Workbooks: vecka_2019.xlsx (")
